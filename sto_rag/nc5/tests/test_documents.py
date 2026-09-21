import io
import tempfile
import unittest
import zipfile
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from nc5.documents import inspect_file,parse,coalesce_groups
from nc5.catalog import load_catalog
from nc5.formatting import measure,check
from unittest.mock import patch

class Documents(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.path=Path(self.tmp.name)/'doc.docx'
    def test_macros_and_entities_rejected(self):
        for name,data in [('word/vbaProject.bin',b'fake'),('word/document.xml',b'<!DOCTYPE x>')]:
            with zipfile.ZipFile(self.path,'w') as z:z.writestr(name,data)
            with self.assertRaises(ValueError):inspect_file(self.path)
    def test_parent_headings_merges_and_real_font(self):
        d=Document();d.add_paragraph('ЧАСТНОЕ ТЕХНИЧЕСКОЕ ЗАДАНИЕ');d.add_heading('1 Требования',level=1);d.add_heading('1.1 Первый подраздел',level=2);d.add_paragraph('Кириллический текст');t=d.add_table(rows=3,cols=3);t.cell(0,0).merge(t.cell(0,1)).text='Группа';t.cell(0,2).text='Количество';t.cell(1,0).text='Позиция';t.cell(1,1).text='Цена';t.cell(1,2).text='Штук';t.cell(2,0).text='A-01';t.cell(2,1).text='18,50';t.cell(2,2).text='3';d.save(self.path)
        result=parse(self.path,load_catalog());self.assertEqual(result['profile']['type'],'ЧТЗ');self.assertEqual(len(result['headings']),2)
        b=next(b for b in result['blocks'] if b['text']=='Группа');self.assertEqual(b['table_context']['span'],2)
        self.assertTrue(any('Первый подраздел' in b['address'] for b in result['blocks']))
        self.assertTrue(measure(self.path));self.assertTrue(coalesce_groups(result))
    def test_deleted_text_excluded(self):
        d=Document();d.add_heading('1 Раздел',level=1);p=d.add_paragraph('Текущий текст. ');deleted=OxmlElement('w:del');r=OxmlElement('w:r');txt=OxmlElement('w:delText');txt.text='Удалённая команда';r.append(txt);deleted.append(r);p._p.append(deleted);d.save(self.path)
        result=parse(self.path,load_catalog());self.assertEqual(result['changes']['deletions'],1);self.assertNotIn('Удалённая команда',''.join(b['text'] for b in result['blocks']))
    def test_annotation_using_toc_heading_style_is_not_excluded(self):
        from docx.enum.style import WD_STYLE_TYPE
        d=Document()
        if 'TOC Heading' not in d.styles:d.styles.add_style('TOC Heading',WD_STYLE_TYPE.PARAGRAPH)
        d.add_paragraph('АННОТАЦИЯ','TOC Heading');d.add_paragraph('Этот документ описывает обмен данными.')
        d.add_paragraph('Содержание','TOC Heading');d.add_heading('1 Раздел',1);d.save(self.path)
        doc=parse(self.path,load_catalog());annotation=next(b for b in doc['blocks'] if b['text']=='АННОТАЦИЯ')
        body=next(b for b in doc['blocks'] if b['text'].startswith('Этот документ'))
        self.assertFalse(annotation['toc']);self.assertTrue(annotation['is_heading']);self.assertEqual(body['section'],annotation['locator'])
        self.assertTrue(next(b for b in doc['blocks'] if b['text']=='Содержание')['toc'])
    def test_unavailable_renderer_is_reported_as_unverified(self):
        from nc5.formatting import check
        d=Document();d.add_paragraph('Учебный текст.');d.save(self.path)
        doc={'path':str(self.path),'sha256':'render-fixture','blocks':[]}
        with patch('nc5.formatting.DATA',Path(self.tmp.name)),patch('nc5.formatting.subprocess.run',side_effect=FileNotFoundError('No renderer')):
            findings,coverage=check(doc,{'cards':[]},'render')
        self.assertFalse(findings)
        self.assertTrue(any(x.get('check')=='render' and x['state']=='unverified' for x in coverage))
    def test_instruction_formatting_checks_page_and_body_properties(self):
        from docx.shared import Inches,Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        d=Document();section=d.sections[0];section.left_margin=Inches(.5);section.right_margin=Inches(.5)
        normal=d.styles['Normal'];normal.font.name='Arial';normal.font.size=Pt(12)
        p=d.add_paragraph('Обычный содержательный абзац.');p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.space_after=Pt(6);d.save(self.path)
        cards=[]
        for clause in ('7.6','27.1','27.4','27.6','27.8'):
            cards.append({'document_name':'Инструкция по делопроизводству','clause':clause,'requirement_id':'instruction-'+clause,'source_quote':'абзацный отступ – 1,25 см' if clause=='7.6' else clause})
        doc={'id':'d','path':str(self.path),'sha256':'fixture','blocks':[{'document':'d','locator':'p1','text':p.text,'heading_path':['Раздел']} ]}
        findings,coverage=check(doc,{'cards':cards},'xml')
        ids={f['requirement_id'] for f in findings}
        self.assertIn('instruction-7.6',ids);self.assertIn('instruction-27.6',ids)
        self.assertTrue(any(x.get('check')=='instruction_page_layout' for x in coverage))
    def test_repeated_formatting_findings_are_compact(self):
        from nc5.formatting import compact_results
        base={'requirement_id':'r','issue':'Неверный интервал','suggestion':'Исправить','kind':'violation','explanation':'Измерено 12 пт'}
        findings=[{**base,'evidence':[{'document':'d','locator':'p'+str(i),'quote':'Текст '+str(i)}]} for i in range(30)]
        result,_=compact_results(findings,[])
        self.assertEqual(len(result),1);self.assertEqual(result[0]['occurrence_count'],30);self.assertEqual(len(result[0]['evidence']),24);self.assertEqual(result[0]['additional_occurrences'],6)

if __name__=='__main__':unittest.main()
