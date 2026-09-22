import tempfile
import unittest
from pathlib import Path
from docx import Document
from nc5.engine import Engine
from nc5.store import Store
from nc5.common import config
from nc5 import feedback
from unittest.mock import patch

class Fake:
    signature='synthetic-test-model';context=24576
    def probe(self):return {'context':self.context,'signature':self.signature,'model':'synthetic-test-model'}
    def count(self,p):return 1000
    def generate(self,p):
        raw={'findings':[],'facts':[],'coverage':[],'decisions':[],'limitations':[]}
        if p['stage']=='language':
            b=next((x for x in p['blocks'] if 'Сведения должна' in x['text']),None)
            if b:raw['findings']=[{'category':'grammar','issue':'Нарушено согласование','kind':'violation','severity':'minor','explanation':'Сведения — множественное число.','suggestion':'Сведения должны храниться.','requirement_id':'','search_query':'','evidence':[{'document':b['document'],'locator':b['locator'],'quote':b['text']}]}]
        if p['stage']=='verify':raw['decisions']=[{'id':c['id'],'verdict':'confirmed','reason':'Проверено согласование и исходная цитата.'} for c in p['candidates']]
        return raw,{'seconds':0,'usage':{},'estimated_tokens':1000}

class EngineTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.root=Path(self.tmp.name)
        runtime=patch('nc5.engine.DATA',self.root/'runtime');runtime.start();self.addCleanup(runtime.stop)
        d=Document();d.add_paragraph('ЧАСТНОЕ ТЕХНИЧЕСКОЕ ЗАДАНИЕ');d.add_heading('1 Сведения',1);d.add_paragraph('Сведения должна храниться.');self.path=self.root/'example.docx';d.save(self.path)
        self.e=Engine(config(),Store(self.root/'db'));self.e.client=Fake()
    def test_end_to_end_queue_and_verification(self):
        j=self.e.create([str(self.path)],{'check_sto':False});self.e.run(j)
        self.assertNotEqual(self.e.store.job(j)['state'],'failed');self.assertTrue(all(t['state']=='done' for t in self.e.store.tasks(j)))
        findings=self.e.store.findings(j);self.assertEqual(len(findings),1);self.assertEqual(findings[0]['status'],'confirmed');self.assertIn('пункт 1 «Сведения»',findings[0]['evidence'][0]['address'])
        self.assertIn('cross',self.e.store.job(j)['data']['derived'])
        second=self.e.create([str(self.path)])
        self.assertTrue(self.e.store.job(second)['data']['options']['check_sto'])
    def test_sto_planning_builds_one_registry_and_internal_evidence_map(self):
        j=self.e.create([str(self.path)],{'check_language':False,'check_logic':False,'check_sto':True})
        self.e.prepare(j);data=self.e.store.job(j)['data'];tasks=self.e.store.tasks(j,'sto')
        self.assertIn(next(iter(data['analysis_registry'])),data['analysis_registry'])
        self.assertTrue((self.root/'runtime'/'jobs'/j/'analysis-registry.json').exists())
        self.assertTrue(tasks)
        self.assertTrue(all('_evidence_map' in task['payload'] for task in tasks))
        # Planner metadata remains available for safe splitting but is absent from
        # the serialized LLM request.
        from nc5.model import Client
        self.assertNotIn('_evidence_map',Client(config()).request(tasks[0]['payload'])['messages'][1]['content'])
    def test_timeout_is_visible(self):
        self.e.client.generate=lambda p:(_ for _ in ()).throw(TimeoutError('test timeout'))
        j=self.e.create([str(self.path)],{'check_sto':False});self.e.run(j)
        self.assertEqual(self.e.store.job(j)['state'],'partial');self.assertTrue(any(t['state']=='failed' for t in self.e.store.tasks(j)))
        errors=self.e.status(j)['task_errors'];self.assertTrue(errors);self.assertIn('test timeout',errors[0]['error']);self.assertIn('stage',errors[0])
    def test_verification_backlog_is_drained_during_section_review(self):
        import copy
        generate=self.e.client.generate
        def three(payload):
            raw,metrics=generate(payload)
            if payload['stage']=='language' and raw['findings']:
                base=raw['findings'][0]
                raw['findings']=[{**copy.deepcopy(base),'suggestion':'Проверенное исправление '+str(i)} for i in range(3)]
            return raw,metrics
        self.e.client.generate=three
        j=self.e.create([str(self.path)],{'check_sto':False});self.e.prepare(j)
        while task:=self.e.store.claim(j,'language'):
            self.e.execute(task)
            if self.e.store.findings(j):break
        self.assertEqual([f['status'] for f in self.e.store.findings(j)],['confirmed']*3)
        self.assertFalse(any(t['state']=='pending' for t in self.e.store.tasks(j,'verify')))
    def test_verified_recommendation_replaces_initial_proposal(self):
        generate=self.e.client.generate
        def corrected(payload):
            raw,metrics=generate(payload)
            if payload['stage']=='verify':
                for d in raw['decisions']:d['suggestion']='Сведения должны храниться в системе.'
            return raw,metrics
        self.e.client.generate=corrected
        j=self.e.create([str(self.path)],{'check_sto':False});self.e.run(j)
        f=self.e.store.findings(j)[0]
        self.assertEqual(f['suggestion'],'Сведения должны храниться в системе.')
        self.assertEqual(f['initial_suggestion'],'Сведения должны храниться.')
    def test_same_mixed_alphabet_defect_is_not_counted_again(self):
        a={'category':'техническая логика','issue':'Кириллица в идентификаторе поля','explanation':'Смешаны алфавиты.','suggestion':'Сверить контракт','evidence':[{'document':'d','locator':'p1','quote':'сountry'}]}
        b={**a,'category':'грамотность','issue':'Опечатка: кириллическая буква','suggestion':'Исправить букву'}
        j=self.e.create([str(self.path)])
        self.assertEqual(self.e.store.finding(j,a),self.e.store.finding(j,b))
        self.assertNotEqual(self.e.store.finding(j,a),self.e.store.finding(j,{**b,'evidence':[{'document':'d','locator':'p2','quote':'сountry'}]}))
    def test_fresh_review_cannot_reuse_previous_job_answers(self):
        a=self.e.create([str(self.path)],{'check_sto':False});self.e.run(a)
        b=self.e.create([str(self.path)],{'check_sto':False,'reuse_cache':False});self.e.run(b)
        old={t['cache_key'] for t in self.e.store.tasks(a)}
        new={t['cache_key'] for t in self.e.store.tasks(b)}
        self.assertFalse(old & new)
        self.assertFalse(any((t['result'] or {}).get('cached') for t in self.e.store.tasks(b)))
        self.assertTrue(all(self.e.store.cached(t['cache_key']) for t in self.e.store.tasks(b)))
    def test_verifier_explanation_is_itself_validated(self):
        generate=self.e.client.generate
        def wrong_reason(payload):
            raw,metrics=generate(payload)
            if payload['stage']=='verify':
                for decision in raw['decisions']:decision['reason']='Это неологизм, поэтому слово ошибочно.'
            return raw,metrics
        self.e.client.generate=wrong_reason
        j=self.e.create([str(self.path)],{'check_sto':False});self.e.run(j)
        findings=self.e.store.findings(j)
        self.assertEqual(len(findings),1);self.assertEqual(findings[0]['status'],'question')
        self.assertIn('источник',findings[0]['verification'])
    def test_bad_group_element_does_not_discard_valid_finding(self):
        import copy
        generate=self.e.client.generate
        def mixed(payload):
            raw,metrics=generate(payload)
            if raw['findings']:
                bad=copy.deepcopy(raw['findings'][0]);bad['evidence'][0]['quote']='Этой строки нет в документе.'
                raw['findings'].append(bad)
            return raw,metrics
        self.e.client.generate=mixed
        j=self.e.create([str(self.path)],{'check_sto':False});self.e.run(j)
        self.assertEqual(self.e.store.job(j)['state'],'partial')
        self.assertEqual([f['status'] for f in self.e.store.findings(j)],['confirmed'])
        self.assertTrue(any((t['result'] or {}).get('invalid') for t in self.e.store.tasks(j)))
    def test_malformed_individual_decision_becomes_question(self):
        generate=self.e.client.generate
        def malformed(payload):
            raw,metrics=generate(payload)
            if payload['stage']=='verify':raw['decisions']=[{'verdict':'confirmed'}]
            return raw,metrics
        self.e.client.generate=malformed
        j=self.e.create([str(self.path)],{'check_sto':False});self.e.run(j)
        self.assertEqual(self.e.store.job(j)['state'],'partial')
        self.assertEqual([f['status'] for f in self.e.store.findings(j)],['question'])
    def test_report_counts_actual_retries_not_queue_claims(self):
        from nc5.report import build
        generate=self.e.client.generate;attempts=[]
        def transient(payload):
            attempts.append(payload['stage'])
            if len(attempts)==1:raise TimeoutError('Temporary test failure')
            return generate(payload)
        self.e.client.generate=transient
        j=self.e.create([str(self.path)],{'check_sto':False});self.e.run(j)
        metrics=build(self.e,j)['metrics']
        self.assertEqual(metrics['retries'],1)
        self.assertEqual(metrics['generation_attempts'],len(attempts))
        self.assertIn('language',metrics['by_stage'])
    def test_resume_refuses_changed_document_without_losing_results(self):
        j=self.e.create([str(self.path)],{'check_sto':False});self.e.run(j)
        original=self.e.store.findings(j)
        d=Document(self.path);d.add_paragraph('Новая редакция входного документа.');d.save(self.path)
        self.e.store.update(j,'paused');self.e.run(j)
        self.assertEqual(self.e.store.job(j)['state'],'failed')
        self.assertIn('Документ изменён',self.e.store.job(j)['data']['fatal_error'])
        self.assertEqual(self.e.store.findings(j),original)
    def test_owner_lesson_isolation(self):
        from nc5.common import dumps
        with self.e.store.connect() as c:c.execute('INSERT INTO lessons VALUES(?,?,?,?)',('case',1,0,dumps({'owner':'alice','type':'ЧТЗ','correction':'test'})))
        d={'profile':{'type':'ЧТЗ'}}
        self.assertEqual(len(self.e.lessons(d,'alice')),1);self.assertEqual(self.e.lessons(d,'bob'),[])

if __name__=='__main__':unittest.main()
