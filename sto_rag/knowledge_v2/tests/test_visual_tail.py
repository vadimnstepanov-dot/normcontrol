import copy,json,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from PIL import Image
from docx import Document
from knowledge_v2.store import KnowledgeStore,Conflict
from knowledge_v2.review import corpus
from knowledge_v2.review_client import LlamaClient
from knowledge_v2 import visual_tail as vt
from knowledge_v2.model_queue import model_turn
from knowledge_v2.model_profile import ensure
from .test_review import ReviewTests,Model


class InventoryTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name);self.store=KnowledgeStore(self.root/'data')
        self.path=self.root/'input.docx';self.image=self.root/'diagram.png'
        Image.new('RGB',(100,80),'white').save(self.image)
        d=Document();d.add_heading('1. Функция',1);d.add_paragraph('Обмен данными')
        d.add_picture(str(self.image));d.add_picture(str(self.image));d.save(self.path)
        self.docs=corpus([self.path]);vt.prepare([self.path],self.docs,self.store)
        self.model=Model()
        self.rows=[dict(id='r',document_id=self.docs[0]['id'],applicability={'result':'applicable'},issues=[],context=[])]

    def test_inventory_deduplicates_and_keeps_occurrences_and_cache(self):
        items=self.docs[0]['visual_inventory']['items'];self.assertEqual(len(items),1)
        self.assertEqual(len(items[0]['occurrences']),2)
        with patch('knowledge_v2.structure_docx.read_docx',side_effect=AssertionError('repeat parse')):
            vt.prepare([self.path],self.docs,self.store)
        self.assertTrue(vt.render(self.store,items[0]).exists())

    def test_planning_keeps_visual_scope_partial_and_no_silent_overflow(self):
        batches,gaps=vt.plan(self.rows,self.docs,self.model)
        self.assertEqual(len(batches),1);self.assertEqual(gaps,[])
        self.assertFalse(batches[0]['payload']['completeness']['full_text'])
        self.model.context=9000
        batches,gaps=vt.plan(self.rows,self.docs,self.model)
        self.assertEqual(batches,[]);self.assertEqual(len(gaps),1)

    def test_no_applicable_norms_has_explicit_gap(self):
        batches,gaps=vt.plan([],self.docs,self.model)
        self.assertEqual(batches,[]);self.assertTrue(gaps)

    def test_modified_asset_refused(self):
        item=self.docs[0]['visual_inventory']['items'][0]
        (self.store.directory/'review-visuals'/item['document_id']/item['original']).write_bytes(b'changed')
        with self.assertRaises(Conflict):vt.render(self.store,item)

    def test_pixel_finding_is_never_confirmed_violation(self):
        batches,_=vt.plan(self.rows,self.docs,self.model);batch=batches[0]
        d=dict(obligation_id='r',outcome='violated',claim='contradiction',reason='Конфликт',observation='Стрелка',bbox=[0,0,1,1],evidence=[])
        findings=vt.findings(batches,{'results':{batch['id']:{'decisions':[d]}}})
        self.assertEqual(findings[0]['state'],'unknown');self.assertTrue(findings[0]['preliminary_violation'])
        self.assertFalse(findings[0]['visual_evidence']['expert_approved'])

    def test_profile_requires_exclusive_turn(self):
        with patch.dict('os.environ',{'KNOWLEDGE_LLM_PROFILE_CONTROL':'1'}):
            with self.assertRaises(RuntimeError):ensure(self.model,'vision')
        with model_turn(self.store,self.model):self.assertTrue(self.model._model_ticket)
        self.assertIsNone(self.model._model_ticket)

    def test_response_validates_alias_coordinates_quotes_and_complete_output(self):
        batches,_=vt.plan(self.rows,self.docs,self.model);batch=dict(batches[0],_render=str(self.image))
        client=LlamaClient.__new__(LlamaClient);client.model='test';client.output_tokens=256
        def reply(path,request):
            alias=request['response_format']['json_schema']['schema']['properties']['decisions']['items']['properties']['obligation_id']['enum'][0]
            value=dict(obligation_id=alias,outcome='unknown',claim='unknown',reason='Неясно',observation='Стрелка не видна',bbox=[0,0,1,1],evidence=[])
            value.update(changes)
            self.assertEqual(request['messages'][1]['content'][1]['type'],'image_url')
            return dict(choices=[dict(finish_reason=finish,message={'content':json.dumps({'decisions':[value]})})])
        client.http=reply;changes={};finish='stop'
        self.assertEqual(vt.complete(client,batch,'visual_check')[0]['obligation_id'],'r')
        for changes in ({'bbox':[1,0,0,1]},{'obligation_id':'foreign'},{'evidence':[{'block_id':'foreign','quote':'invented'}]}):
            with self.assertRaises(ValueError):vt.complete(client,batch,'visual_check')
        changes={};finish='length'
        with self.assertRaises(ValueError):vt.complete(client,batch,'visual_check')


class CursorTests(ReviewTests):
    # Reuse canonical release fixture, not live data or model requests.
    def test_visual_pause_resume_switches_once_per_tail_and_keeps_results(self):
        parent=self.create()
        with self.store.connection() as db:payload=json.loads(db.execute('SELECT payload FROM tasks WHERE id=?',(parent,)).fetchone()[0])
        image=dict(document_id=self.did,locator='p2',sha256='b'*64,occurrences=['p2'])
        batches=[dict(id=str(i),image=image,payload={'obligations':payload['rows']}) for i in range(2)]
        task=self.store.enqueue('visual.run','test-visual',dict(batches=batches,snapshot_id=payload['snapshot_id'],model_signature=self.model.signature))
        self.model.http=lambda path:{'modalities':{'vision':True}}
        reply=[dict(obligation_id=payload['rows'][0]['id'],outcome='unknown',reason='Не читается',observation='',bbox=[0,0,1,1],evidence=[])]
        with patch('knowledge_v2.model_profile.ensure') as switch,patch.object(vt,'render',return_value=self.path),patch.object(vt,'complete',return_value=reply) as call:
            result=vt.run(self.store,task,self.model,lambda sid:True,lambda done,total,result:done==1)
            self.assertEqual(result['state'],'paused');self.assertEqual(call.call_count,1)
            self.assertEqual([x.args[1] for x in switch.call_args_list],['vision','text'])
            result=vt.run(self.store,task,self.model,lambda sid:True,lambda *args:False)
            self.assertEqual(result['state'],'done');self.assertEqual(call.call_count,2)
            self.assertEqual(len(result['results']),2)

    def test_visual_failure_restores_text_and_is_durable(self):
        parent=self.create()
        with self.store.connection() as db:payload=json.loads(db.execute('SELECT payload FROM tasks WHERE id=?',(parent,)).fetchone()[0])
        task=self.store.enqueue('visual.run','test-failure',dict(batches=[dict(id='x',image={})],snapshot_id=payload['snapshot_id'],model_signature=self.model.signature))
        self.model.http=lambda path:{'modalities':{'vision':True}}
        with patch('knowledge_v2.model_profile.ensure') as switch,patch.object(vt,'render',side_effect=ValueError('unreadable')):
            result=vt.run(self.store,task,self.model,lambda sid:True,lambda *args:False)
            self.assertEqual(result['failures']['x']['error'],'ValueError')
            self.assertEqual(switch.call_args_list[-1].args[1],'text')
