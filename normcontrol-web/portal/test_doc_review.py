import base64,tempfile
from pathlib import Path
from unittest import TestCase
from unittest.mock import patch,Mock
from zipfile import ZipFile
from docx import Document
from lxml import etree as E
from word_source import VERSION,OLE,parts,anchor_map,PKG
from . import word_review as shared,doc_review

class NativeDocTests(TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.source=Path(self.tmp.name)/'source.doc';self.source.write_bytes(OLE+b'private synthetic source')
        x=Path(self.tmp.name)/'same.docx';d=Document();d.add_paragraph('😀 Проверка содержит ошибка.');d.add_paragraph('Повтор');d.add_paragraph('Повтор');d.save(x)
        root=E.Element(PKG+'package',nsmap={'pkg':PKG[1:-1]})
        with ZipFile(x) as z:
            for name in z.namelist():
                p=E.SubElement(root,PKG+'part',{PKG+'name':'/'+name})
                if name.endswith(('.xml','.rels')):E.SubElement(p,PKG+'xmlData').append(shared.xml(z.read(name)))
                else:E.SubElement(p,PKG+'binaryData').text=base64.b64encode(z.read(name)).decode()
        self.snapshot={'adapter':VERSION,'source_sha256':shared.sha(self.source),'flat_xml':E.tostring(root).decode(),'paragraphs':[]}
        offset=0
        for value in ('😀 Проверка содержит ошибка.','Повтор','Повтор'):
            length=len((value+'\r').encode('utf-16-le'))//2
            self.snapshot['paragraphs'].append({'text':value,'raw':value+'\r','start':offset,'end':offset+length,'editable':True});offset+=length
        self.description={'id':1,'aliases':['target'],'source_sha256':shared.sha(self.source),'working_sha256':shared.sha(self.source),'role':'target'}
        self.finding={'id':'F1','status':'confirmed','issue':'Ошибка','evidence':[{'document':'target','locator':'p1','quote':'ошибка'}]}
        self.client=Mock();self.client.read.return_value=self.snapshot
        p=patch.object(doc_review,'client',return_value=self.client);p.start();self.addCleanup(p.stop)
    def draft(self,**kw):return doc_review.plan(self.source,self.source,self.description,'v1',[self.finding],**kw)
    def proposal(self,draft):
        op=draft['operations'][0]
        return {**{k:op[k] for k in ('document_id','result_version','source_sha256','working_sha256','part','locator','start','end','original')},'type':'replace','proposed':'исправление','verification':'specialist_confirmed'}
    def test_same_shared_plan_and_utf16_native_range(self):
        draft=self.draft();saved=self.draft(proposals={'F1':self.proposal(draft)});op=saved['operations'][0]
        self.assertEqual(saved['output_format'],'doc');self.assertEqual(op['type'],'replace')
        self.assertEqual(op['native_start'],op['start']+1);self.assertIn('Основание:',op['comment'])
        self.assertEqual(op['original'],'ошибка');self.assertFalse(list(Path(self.tmp.name).glob('*.tmp.docx')))
    def test_existing_revision_never_replaced(self):
        draft=self.draft();self.snapshot['paragraphs'][0]['editable']=False
        saved=self.draft(proposals={'F1':self.proposal(draft)});op=saved['operations'][0]
        self.assertEqual(op['type'],'comment');self.assertEqual(op['native_start'],0);self.assertEqual(op['native_end'],self.snapshot['paragraphs'][0]['end'])
    def test_only_selected_native_ranges_are_requested(self):
        row=self.snapshot['paragraphs'][0]
        self.snapshot.update(paragraph_count=3,paragraphs=[])
        self.client.anchors.return_value={'p1':row}
        saved=self.draft()
        request=self.client.anchors.call_args.args
        self.assertEqual(len(request[1]),1);self.assertEqual(request[1][0]['locator'],'p1');self.assertEqual(request[1][0]['paragraph'],1)
        self.assertEqual(len(saved['operations']),1)
    def test_unproven_repeat_mapping_skipped(self):
        root=shared.xml(parts(self.snapshot)['word/document.xml']);ps=shared.paragraphs(root)
        self.snapshot['paragraphs'].pop()
        self.assertNotIn('p2',anchor_map(self.snapshot,ps,shared.text));self.assertIn('p1',anchor_map(self.snapshot,ps,shared.text))
    def test_source_hash_mismatch_rejected(self):
        self.description['source_sha256']='changed'
        self.assertRaises(shared.ReviewError,self.draft)
    def test_unsafe_flat_package_rejected(self):
        bad=dict(self.snapshot,flat_xml='<!DOCTYPE a><a/>')
        self.assertRaises(ValueError,parts,bad)
    def test_unknown_adapter_cannot_write_doc(self):
        self.assertRaises(shared.ReviewError,doc_review.generate,self.source,self.source,{},Path(self.tmp.name)/'out.doc')
