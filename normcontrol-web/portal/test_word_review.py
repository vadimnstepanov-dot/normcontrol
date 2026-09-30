import copy,hashlib,json,tempfile,uuid
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from io import BytesIO
from unittest import TestCase as UnitTestCase
from docx import Document as Docx
from django.test import TestCase,override_settings
from django.contrib.auth.models import User
from django.core.files.base import ContentFile
from . import word_review as wr
from .models import Batch,Document,WorkerRun,WordReviewExport,FindingDisposition

def fixture(path):
 d=Docx();d.sections[0].header.paragraphs[0].text='Колонтитул'
 p=d.add_paragraph();p.add_run('Начало ');p.add_run('оши').bold=True;p.add_run('бка & <текст>');p.add_run('\tконец')
 d.add_paragraph('Дубль дубль.');d.add_paragraph('Поле 1');d.add_paragraph('Прежняя правка')
 table=d.add_table(rows=2,cols=2);table.cell(0,0).merge(table.cell(0,1));table.cell(0,0).text='Объединено';table.cell(1,0).text='Ячейка ошибка';table.cell(1,1).text='Формула'
 d.add_paragraph('Закладка и гиперссылка');d.save(path)
 with ZipFile(path) as z:parts={i.filename:z.read(i) for i in z.infolist()}
 root=wr.xml(parts['word/document.xml']);ps=wr.paragraphs(root)
 r=ps[2].find(wr.Q+'r');r.append(wr.node('fldChar',fldCharType='begin'));r.append(wr.node('instrText'));r[-1].text=' REF target ';r.append(wr.node('fldChar',fldCharType='end'))
 old=wr.node('ins',id='900',author='Прежний автор',date='2020-01-01T00:00:00Z');old.append(wr.run(wr.node('r'),' ранее'));ps[3].append(old)
 ps[-1].insert(0,wr.node('bookmarkStart',id='8',name='target'));ps[-1].append(wr.node('bookmarkEnd',id='8'))
 parts['word/document.xml']=wr.dump(root)
 with ZipFile(path,'w',ZIP_DEFLATED) as z:
  for n,v in parts.items():z.writestr(n,v)

def finding(fid='F-1',quote='ошибка',locator='p1',status='confirmed'):
 return {'id':fid,'status':status,'issue':'Ошибка в тексте','explanation':'Проверена конкретная цитата','suggestion':'Уточнить редакцию','evidence':[{'document':'target','locator':locator,'quote':quote}]}

def view_text(root,ids,accept):
 result=copy.deepcopy(root)
 for n in list(result.iter()):
  if n.tag not in (wr.Q+'ins',wr.Q+'del') or n.get(wr.Q+'id') not in ids:continue
  parent=n.getparent();index=parent.index(n)
  keep=(n.tag==wr.Q+'ins')==accept
  if keep:
   for c in list(n):
    for t in c.iter(wr.Q+'delText'):t.tag=wr.Q+'t'
    parent.insert(index,c);index+=1
  parent.remove(n)
 return result

class WordReviewUnitTests(UnitTestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.source=Path(self.tmp.name)/'source.docx';fixture(self.source)
  self.document={'id':1,'source_sha256':wr.sha(self.source),'working_sha256':wr.sha(self.source),'aliases':['target'],'reference_aliases':['ref'],'role':'target'}
 def draft(self,fs,**kw):return wr.plan(self.source,self.source,self.document,'version-a',fs,**kw)
 def approved(self,draft,new,kind='replace'):
  op=draft['operations'][0];return {**{k:op[k] for k in ('document_id','result_version','source_sha256','working_sha256','part','locator','start','end','original')},'type':kind,'proposed':new,'verification':'specialist_confirmed','basis':'Проверено специалистом'}
 def generate(self,p):
  out=Path(self.tmp.name)/'out.docx';r=wr.generate(self.source,self.source,p,out);return out,r,wr.load(out)
 def test_multirun_replace_reject_accept_and_unchanged_parts(self):
  fs=[finding()];draft=self.draft(fs);p=self.draft(fs,proposals={'F-1':self.approved(draft,'исправление')});before=wr.sha(self.source)
  out,r,root=self.generate(p);self.assertEqual(r['edits'],1);self.assertEqual(r['comments'],1)
  original=wr.load(self.source);expected=wr.text(wr.paragraphs(original)[0]).replace('ошибка','исправление')
  self.assertEqual(wr.text(wr.paragraphs(view_text(root,r['new_revision_ids'],True))[0]),expected)
  self.assertEqual([wr.text(x) for x in wr.paragraphs(view_text(root,r['new_revision_ids'],False))],[wr.text(x) for x in wr.paragraphs(original)])
  self.assertEqual(wr.dump(root.xpath('.//w:ins[@w:id="900"]',namespaces=wr.NS)[0]),wr.dump(original.xpath('.//w:ins[@w:id="900"]',namespaces=wr.NS)[0]))
  with ZipFile(self.source) as a,ZipFile(out) as b:
   for name in a.namelist():
    if name not in ('word/document.xml','word/settings.xml','word/_rels/document.xml.rels','[Content_Types].xml','docProps/core.xml'):self.assertEqual(a.read(name),b.read(name),name)
  self.assertEqual(wr.sha(self.source),before)
 def test_insert_delete_xml_symbols_and_tabs(self):
  fs=[finding('insert','Начало'),finding('delete','<текст>')];a=self.draft([fs[0]]);b=self.draft([fs[1]])
  p=self.draft(fs,proposals={'insert':self.approved(a,' & новое','insert'),'delete':self.approved(b,'','delete')});out,r,root=self.generate(p)
  value=wr.text(wr.paragraphs(view_text(root,r['new_revision_ids'],True))[0]);self.assertIn('Начало & новое',value);self.assertNotIn('<текст>',value);self.assertIn('\tконец',value);self.assertEqual(r['edits'],2)
 def test_duplicate_quote_never_edits(self):
  f=finding(quote='дубль',locator='p2');f['evidence'][0]['quote']='Дубль';# case-sensitive first quote is unique
  f['evidence'][0]['quote']='ль';d=self.draft([f]);self.assertIn('неоднозначен',d['operations'][0]['reason']);_,r,_=self.generate(d);self.assertEqual(r['edits'],0)
 def test_merge_identical_edits_and_conflicts_are_comments(self):
  f=finding();g=finding('F-2');a=self.draft([f]);b=self.draft([g]);p=self.draft([f,g],proposals={'F-1':self.approved(a,'верно'),'F-2':self.approved(b,'верно')})
  self.assertEqual(p['operations'][0]['finding_ids'],['F-1','F-2']);self.assertEqual(self.generate(p)[1]['edits'],1)
  p=self.draft([f,g],proposals={'F-1':self.approved(a,'верно'),'F-2':self.approved(b,'иначе')});self.assertTrue(all(o['type']=='comment' for o in p['operations']))
 def test_fields_old_revisions_and_nested_objects_are_not_edited(self):
  for f in (finding(quote='1',locator='p3'),finding(quote='ранее',locator='p4')):
   draft=self.draft([f]);p=self.draft([f],proposals={'F-1':self.approved(draft,'2')});out,r,root=self.generate(p);self.assertEqual(r['edits'],0);self.assertEqual(r['comments'],1)
 def test_status_decision_and_missing_range_outcomes(self):
  fs=[finding('fixed'),finding('disputed'),finding('candidate',status='candidate'),finding('rejected',status='rejected'),finding('question',status='question'),finding('bad',locator='header:p1')]
  p=self.draft(fs,decisions={'fixed':{'state':'fixed'},'disputed':{'state':'disputed'}});self.assertEqual(len(p['outcomes']),5);self.assertEqual(len(p['operations']),1)
  p=self.draft([finding(status='candidate')],include_preliminary=True);self.assertIn('гипотеза',p['operations'][0]['comment'])
 def test_changed_hash_invalidates_plan(self):
  p=self.draft([finding()]);self.source.write_bytes(self.source.read_bytes()+b'changed')
  with self.assertRaises(wr.ReviewError):self.generate(p)
 def test_existing_comments_and_modern_extensions_are_preserved(self):
  with ZipFile(self.source) as z:parts={n:z.read(n) for n in z.namelist()}
  root=wr.xml(parts['word/document.xml']);p=wr.paragraphs(root)[0]
  p.insert(0,wr.node('commentRangeStart',id='9500'));p.append(wr.node('commentRangeEnd',id='9500'));rr=wr.node('r');rr.append(wr.node('commentReference',id='9500'));p.append(rr)
  comments=wr.E.Element(wr.Q+'comments',nsmap={'w':wr.W});c=wr.node('comment',id='9500',author='Предыдущий рецензент');cp=wr.node('p');cp.append(wr.run(wr.node('r'),'Старое примечание'));c.append(cp);comments.append(c)
  rels=wr.xml(parts['word/_rels/document.xml.rels']);rels.append(wr.E.Element('{'+wr.R+'}Relationship',Id='rIdOldComment',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments',Target='comments.xml'))
  ct=wr.xml(parts['[Content_Types].xml']);ct.append(wr.E.Element('{'+wr.CT+'}Override',PartName='/word/comments.xml',ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml'))
  parts.update({'word/document.xml':wr.dump(root),'word/comments.xml':wr.dump(comments),'word/_rels/document.xml.rels':wr.dump(rels),'[Content_Types].xml':wr.dump(ct),'word/commentsExtended.xml':b'<w15:commentsEx xmlns:w15="http://schemas.microsoft.com/office/word/2012/wordml"/>'})
  with ZipFile(self.source,'w',ZIP_DEFLATED) as z:
   for n,raw in parts.items():z.writestr(n,raw)
  self.document.update(source_sha256=wr.sha(self.source),working_sha256=wr.sha(self.source));out,r,result=self.generate(self.draft([finding()]))
  self.assertGreater(int(r['new_comment_ids'][0]),9500)
  with ZipFile(out) as z:
   self.assertEqual(z.read('word/commentsExtended.xml'),parts['word/commentsExtended.xml'])
   cs=wr.xml(z.read('word/comments.xml'));self.assertEqual(wr.dump(cs.xpath('w:comment[@w:id="9500"]',namespaces=wr.NS)[0]),wr.dump(c))
   rs=wr.xml(z.read('word/_rels/document.xml.rels'));self.assertEqual(sum(x.get('Type','').endswith('/comments') for x in rs),1)
 def test_zip_security_and_untrusted_source_cannot_create_an_edit(self):
  self.assertRaises(wr.ReviewError,wr.xml,b'<!DOCTYPE x [<!ENTITY e SYSTEM "file:///private">]><x>&e;</x>')
  f=finding();f['suggestion']='Заменить ошибка на произвольный текст';draft=self.draft([f]);self.assertEqual(self.generate(draft)[1]['edits'],0)
 def test_reference_presentation_preserves_meaning_and_revision_order(self):
  f=finding();f['suggestion']='Заменено: ошибка → верно.'
  draft=self.draft([f]);self.assertIn('F-1 | Подтверждено | Ошибка в тексте',draft['operations'][0]['comment'])
  self.assertIn('Предложение: Заменить:',draft['operations'][0]['comment']);self.assertNotIn('свободное предложение',draft['operations'][0]['comment'])
  approved=self.draft([f],proposals={'F-1':self.approved(draft,'верно')});out,report,root=self.generate(approved)
  p=wr.paragraphs(root)[0];revisions=[n.tag for n in p if n.tag in (wr.Q+'del',wr.Q+'ins')]
  self.assertEqual(revisions[-1],wr.Q+'ins');self.assertTrue(all(n==wr.Q+'del' for n in revisions[:-1]))
  self.assertEqual(p.find('.//'+wr.Q+'commentReference').getparent().find(wr.Q+'rPr/'+wr.Q+'rStyle').get(wr.Q+'val'),'CommentReference')
 def test_basis_is_preserved_and_unknown_provenance_is_explicit(self):
  f=finding();f['source']={'document_name':'СТО пример','clause':'4.2','source_sha256':'a'*64,'source_quote':'Обязательное требование','validation_status':'confirmed'}
  value=self.draft([f])['operations'][0]['comment'];self.assertIn('Основание: СТО пример; 4.2; «Обязательное требование»',value)
  del f['source']['validation_status'];value=self.draft([f])['operations'][0]['comment'];self.assertIn('Обязательное требование',value);self.assertIn('подтверждение источника не передано',value)
  del f['source'];value=self.draft([f])['operations'][0]['comment'];self.assertIn('Основание: исходный текст',value);self.assertIn('Отдельное нормативное основание',value)
 def test_hyperlink_and_formula_fallback_preserves_objects(self):
  with ZipFile(self.source) as z:parts={n:z.read(n) for n in z.namelist()}
  root=wr.xml(parts['word/document.xml']);p=wr.paragraphs(root)[0]
  link=wr.node('hyperlink',anchor='target');link.append(wr.run(wr.node('r'),'Ссылка'));p.append(link)
  equation=wr.E.Element('{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath');equation.text='';p.append(equation)
  parts['word/document.xml']=wr.dump(root)
  with ZipFile(self.source,'w',ZIP_DEFLATED) as z:
   for n,raw in parts.items():z.writestr(n,raw)
  self.document.update(source_sha256=wr.sha(self.source),working_sha256=wr.sha(self.source));f=finding(quote='Ссылка');draft=self.draft([f]);out,r,result=self.generate(self.draft([f],proposals={'F-1':self.approved(draft,'Другая')}))
  self.assertEqual(r['edits'],0);self.assertEqual(wr.dump(result.find('.//'+wr.Q+'hyperlink')),wr.dump(link));self.assertEqual(wr.E.tostring(result.find('.//'+equation.tag),method='c14n'),wr.E.tostring(equation,method='c14n'))
 def test_whitespace_map_and_table_anchor(self):
  f=finding(quote='<текст> конец');d=self.draft([f]);self.assertEqual(d['operations'][0]['original'],'<текст>\tконец')
  root=wr.load(self.source);index=next(i for i,p in enumerate(wr.paragraphs(root),1) if wr.text(p)=='Ячейка ошибка')
  f=finding(locator='p'+str(index));draft=self.draft([f]);p=self.draft([f],proposals={'F-1':self.approved(draft,'верно')});self.assertEqual(self.generate(p)[1]['edits'],1)
 def test_missing_optional_settings_and_core_are_created(self):
  with ZipFile(self.source) as z:parts={n:z.read(n) for n in z.namelist() if n not in ('word/settings.xml','docProps/core.xml')}
  for name in ('word/_rels/document.xml.rels','_rels/.rels'):
   root=wr.xml(parts[name])
   for n in list(root):
    if n.get('Type','').endswith(('/settings','/core-properties')):root.remove(n)
   parts[name]=wr.dump(root)
  root=wr.xml(parts['[Content_Types].xml'])
  for n in list(root):
   if n.get('PartName') in ('/word/settings.xml','/docProps/core.xml'):root.remove(n)
  parts['[Content_Types].xml']=wr.dump(root)
  with ZipFile(self.source,'w',ZIP_DEFLATED) as z:
   for n,raw in parts.items():z.writestr(n,raw)
  self.document.update(source_sha256=wr.sha(self.source),working_sha256=wr.sha(self.source));out,_,_=self.generate(self.draft([finding()]))
  with ZipFile(out) as z:
   self.assertEqual(wr.xml(z.read('word/settings.xml')).find(wr.Q+'trackRevisions').get(wr.Q+'val'),'true')
   self.assertEqual(wr.xml(z.read('docProps/core.xml')).find('{http://purl.org/dc/elements/1.1/}creator').text,'NormControl')

@override_settings(KNOWLEDGE_V2_ENABLED=False,WORD_REVIEW_BACKGROUND=False)
class WordReviewWebTests(TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup);self.settings_override=override_settings(MEDIA_ROOT=self.tmp.name);self.settings_override.enable();self.addCleanup(self.settings_override.disable)
  self.user=User.objects.create_user('reviewer');self.other=User.objects.create_user('other-reviewer');self.client.force_login(self.user)
  self.batch=Batch.objects.create(owner=self.user,name='Комплект',status='running');path=Path(self.tmp.name)/'fixture.docx';fixture(path);content=path.read_bytes()
  self.document=Document.objects.create(batch=self.batch,name='Документ.docx',file=ContentFile(content,name='fixture.docx'),size=len(content),sha256=hashlib.sha256(content).hexdigest(),review_role='target')
  self.reference=Document.objects.create(batch=self.batch,name='Эталон.docx',file=ContentFile(content,name='ref.docx'),size=len(content),sha256=hashlib.sha256(content).hexdigest(),review_role='approved_reference')
  fs=[finding('F-'+str(i),status='question') for i in range(65)]
  for f in fs:f['evidence'][0]['document']=self.document.sha256[:20]
  self.run=WorkerRun.objects.create(batch=self.batch,state='running',report={'findings':fs},sequence=1)
  self.url='/normcontol/batches/'+str(self.batch.pk)+'/word-review/'
 def test_full_register_snapshot_and_cached_generation(self):
  response=self.client.post(self.url,{'action':'plan','documents':[self.document.pk],'page':'2','status':'rejected'});self.assertEqual(response.status_code,302)
  export=WordReviewExport.objects.get(batch=self.batch);self.assertEqual(len(export.plans[0]['operations']),65)
  url=self.url+str(export.pk)+'/file/';self.assertEqual(self.client.post(url).status_code,302);export.refresh_from_db();self.assertEqual(export.result['comments'],65);self.assertEqual(export.result['edits'],0)
  original=export.artifact.read();export.artifact.close();self.assertEqual(self.client.post(url).status_code,302);export.refresh_from_db();self.assertEqual(export.artifact.read(),original);export.artifact.close()
  response=self.client.get(url);self.assertEqual(response.status_code,200);response.close();self.assertFalse(FindingDisposition.objects.exists())
 def test_reference_wrong_owner_and_changed_source(self):
  self.assertEqual(self.client.post(self.url,{'action':'plan','documents':[self.reference.pk]}).status_code,200);self.assertFalse(WordReviewExport.objects.exists())
  self.client.post(self.url,{'action':'plan','documents':[self.document.pk]});export=WordReviewExport.objects.get();url=self.url+str(export.pk)+'/file/'
  self.client.force_login(self.other);self.assertEqual(self.client.get(url).status_code,404);self.assertEqual(self.client.post(url).status_code,404)
  self.client.force_login(self.user);Path(self.document.file.path).write_bytes(b'changed');self.assertEqual(self.client.post(url).status_code,409);export.refresh_from_db();self.assertEqual(export.state,'failed')
 def test_intermediate_and_explicit_roles_in_interface(self):
  response=self.client.get(self.url);self.assertContains(response,'Промежуточный результат');self.assertContains(response,'Утверждённый эталон')
  self.client.post(self.url,{'action':'roles','role_'+str(self.document.pk):'approved_reference','role_'+str(self.reference.pk):'target'});self.document.refresh_from_db();self.assertEqual(self.document.review_role,'approved_reference')
 def test_same_plan_is_idempotent_and_building_does_not_generate_twice(self):
  payload={'action':'plan','documents':[self.document.pk]}
  self.client.post(self.url,payload);self.client.post(self.url,payload);self.assertEqual(WordReviewExport.objects.count(),1)
  record=WordReviewExport.objects.get();record.state='building';record.save();response=self.client.post(self.url+str(record.pk)+'/file/');self.assertEqual(response.status_code,409);record.refresh_from_db();self.assertFalse(record.artifact)
 @override_settings(WORD_REVIEW_BACKGROUND=True)
 def test_background_queue_uses_saved_plan_and_finishes(self):
  from unittest.mock import patch
  from .review_export_queue import drain
  self.client.post(self.url,{'action':'plan','documents':[self.document.pk]});record=WordReviewExport.objects.get()
  with patch('portal.review_export_queue.kick') as kick:
   self.assertEqual(self.client.post(self.url+str(record.pk)+'/file/').status_code,302);kick.assert_called_once()
  record.refresh_from_db();self.assertEqual(record.state,'queued');drain();record.refresh_from_db();self.assertEqual(record.state,'ready')
 def test_multiple_targets_are_separate_docx_in_zip(self):
  self.reference.review_role='target';self.reference.save();self.client.post(self.url,{'action':'plan','documents':[self.document.pk,self.reference.pk]})
  record=WordReviewExport.objects.get();response=self.client.post(self.url+str(record.pk)+'/file/');self.assertEqual(response.status_code,302);record.refresh_from_db()
  with record.artifact.open('rb') as handle:
   with ZipFile(handle) as z:self.assertEqual(sum(n.endswith('.docx') for n in z.namelist()),2);self.assertIn('Применение.json',z.namelist())
 def test_reference_change_invalidates_cached_download(self):
  self.client.post(self.url,{'action':'plan','documents':[self.document.pk]});record=WordReviewExport.objects.get();url=self.url+str(record.pk)+'/file/'
  self.client.post(url);Path(self.reference.file.path).write_bytes(b'changed-reference')
  self.assertEqual(self.client.get(url).status_code,409)
 def test_canonical_worker_upload_requires_auth_batch_and_hash(self):
  import base64,os
  from unittest.mock import patch
  token='synthetic-working-copy-token-1234567890';content=Path(self.document.file.path).read_bytes()
  url='/normcontol/worker/'+str(self.run.lease)+'/review-source/'+str(self.document.pk)+'/'
  payload={'source_sha256':self.document.sha256,'working_sha256':hashlib.sha256(content).hexdigest(),'docx':base64.b64encode(content).decode()}
  with patch.dict(os.environ,{'NORMCONTROL_WORKER_TOKEN':token}):
   self.assertEqual(self.client.post(url,json.dumps(payload),content_type='application/json').status_code,403)
   payload['working_sha256']='changed'
   self.assertEqual(self.client.post(url,json.dumps(payload),content_type='application/json',HTTP_AUTHORIZATION='Bearer '+token).status_code,400)
   payload['working_sha256']=hashlib.sha256(content).hexdigest()
   self.assertEqual(self.client.post(url,json.dumps(payload),content_type='application/json',HTTP_AUTHORIZATION='Bearer '+token).status_code,200)
   self.document.refresh_from_db();self.assertEqual(self.document.review_working_sha256,payload['working_sha256'])
   foreign=Document.objects.create(batch=Batch.objects.create(owner=self.other,name='Чужой'),name='foreign.doc',file=ContentFile(content,name='foreign.doc'),size=len(content),sha256=self.document.sha256)
   url='/normcontol/worker/'+str(self.run.lease)+'/review-source/'+str(foreign.pk)+'/'
   self.assertEqual(self.client.post(url,json.dumps(payload),content_type='application/json',HTTP_AUTHORIZATION='Bearer '+token).status_code,404)
 @override_settings(KNOWLEDGE_V2_ENABLED=True)
 def test_normative_snapshot_keeps_basis_citation(self):
  from unittest.mock import patch,MagicMock
  from types import SimpleNamespace
  from .review_export_views import capture
  decision={'state':'violated','obligation':{'id':'O-1','source':{'filename':'СТО тест','sha256':'a'*64},'atom':{'description':'Требование','citation':{'locator':'p754','quote':'Точная нормативная цитата','source_sha256':'a'*64}}},'evidence':[]}
  chunks=MagicMock();chunks.order_by.return_value=[SimpleNamespace(entries=[decision])]
  job=SimpleNamespace(pk=uuid.uuid4(),state='completed',progress={},summary={},finding_chunks=chunks)
  with patch('knowledge.models.KnowledgeCheck.objects.filter') as query,patch('knowledge.checks.visible'):
   query.return_value.select_related.return_value=[job];snapshot,_=capture(self.user,self.batch)
  item=next(f for f in snapshot['findings'] if f['id']=='O-1')
  self.assertEqual(item['source']['source_quote'],'Точная нормативная цитата');self.assertEqual(item['source']['document_name'],'СТО тест');self.assertEqual(item['source']['source_sha256'],'a'*64)
